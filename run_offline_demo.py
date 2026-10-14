# -*- coding: utf-8 -*-
"""法元 · Faron Agent Group —— 离线端到端演示（对齐真实实现）。

不依赖任何模型接口，完整走通一次团队协同，并产出真实交付物：

    ① 部署团队（6 成员 / 12 工具）
    ② 图谱策展：材料 → 案由 → 骨架（要件 / 请求权基础 / 证据 / 争点 / 文书）
    ③ 检索员：真实检索日美法条，登记「工具真实返回」
    ④ 起草员：依模板起草，**故意埋入 1 处库外条号**
    ⑤ 质检员：**两段自检** —— 引用真实性 + 法条适用性
    ⑥ 导出 .docx（含 Word 原生表格）

【为什么做「离线演示」】
    openJiuwen 的 TeamAgent 负责编排，但要驱动 LLM 成员需要模型 Key。
    本脚本把**不依赖模型的那部分**（工具 / 图谱 / 两段自检 / 交付）完整跑出来，
    使得：① 评审无需配置任何 Key 即可复现核心结论；
          ② 「确定性自检」的可验证性得到证明——它本来就不该依赖模型。
    需要真模型时用 `run_online.py`。
"""

from __future__ import annotations

import asyncio
import json
import os
import sys
import time
from pathlib import Path

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "src"))

from faron_agentgroup import (
    KG,
    REGISTRY,
    TASK_KINDS,
    FaronAgentGroup,
    applicability_report,
    audit_applicability,
    audit_citations,
    audit_report,
    kind_meta,
    laws_search,
    template_load,
    kg_find_case,
    kg_query,
    kg_add_candidate,
)

OUT = Path("runtime")
OUT.mkdir(parents=True, exist_ok=True)
TS = time.strftime("%Y%m%d-%H%M%S")
RUNLOG = OUT / f"runlog_{TS}.json"
DELIVERABLE = OUT / f"deliverable_{TS}.docx"

TASK_KIND = "review"
COMMISSION = (
    "我方（中国某制造企业）与日本某商社设立合资公司，合资合同履行中出现出资迟延与"
    "经营分歧。请就下列问题出具《合同审查意见书》：①日方出资迟延是否构成违约、"
    "我方能否请求继续履行；②准据法与管辖约定是否有效；③风险条款与修改建议。"
)

log: dict = {"started_at": time.strftime("%Y-%m-%dT%H:%M:%S"),
             "task_kind": TASK_KIND, "kind_label": kind_meta(TASK_KIND)["label"],
             "commission": COMMISSION, "steps": []}


def step(name: str, payload) -> None:
    log["steps"].append({"step": name, "at": time.strftime("%H:%M:%S"), "payload": payload})
    print(f"\n{'-' * 76}\n> {name}\n{'-' * 76}")
    print(payload if isinstance(payload, str)
          else json.dumps(payload, ensure_ascii=False, indent=2))


async def call(tool_obj, inputs: dict):
    """调用一个 LocalFunction 工具并解析 JSON 返回。"""
    raw = await tool_obj.invoke(inputs)
    if isinstance(raw, str):
        return json.loads(raw)
    if hasattr(raw, "data"):
        d = raw.data
        if isinstance(d, str):
            return json.loads(d)
        if isinstance(d, dict) and "content" in d:
            c = d["content"]
            return json.loads(c) if isinstance(c, str) else c
        return d
    return raw


def add_candidate_direct(case_name: str, law_title: str):
    """同步登记候选边（底层，跳过异步工具包装，便于演示）。"""
    return KG.add_candidate(case_name, law_title)


async def main() -> int:
    print("=" * 76)
    print("法元 · Faron Agent Group —— 离线端到端演示（对齐真实实现）")
    print("=" * 76)

    # ---------------- ① 部署团队 ----------------
    group = FaronAgentGroup()
    roster = group.describe()
    step("STEP 1 · 部署 Agent Group（openJiuwen agent_teams）", roster)
    print(f"\n团队已构建：TeamAgent role={group.agent.role.value}"
          f" / 已注册工具 {group.tools_registered} 个")

    REGISTRY.clear()

    # ---------------- ② 图谱策展：定案由 + 取骨架 ----------------
    case_resp = await call(kg_find_case, {"text": COMMISSION})
    case_name = case_resp["case"]
    sk_resp = await call(kg_query, {"case_name": case_name})
    sk = sk_resp["skeleton"]
    step("STEP 2 · 图谱策展（kg_find_case → kg_query）", {
        "命中案由": case_name,
        "骨架渲染": sk_resp["described"],
        "请求权基础（法律名）": sk["lawNames"],
        "证据种类数": len({e for el in sk["elements"] for e in el["evidences"]}),
        "注意": "图谱法条节点只有法律名、没有条号——条号必须由检索员真实检索",
    })

    # ---------------- ③ 检索员：真实检索 ----------------
    searches = [
        ("会社法", "JP", "日本公司法（图谱请求权基础）"),
        ("契約", "JP", "日本民法（合同履行）"),
        ("trademark", "US", "美国法典 Title 15（图谱证据基础）"),
    ]
    search_log = []
    for kw, country, why in searches:
        r = await call(laws_search, {"keyword": kw, "country": country})
        search_log.append({
            "关键词": kw, "国家": country, "用途": why,
            "命中": [f'{h["title"]} {h["article_no"]}' for h in r["hits"]],
        })
    step("STEP 3 · 检索员真实检索（laws_search）", {
        "检索明细": search_log,
        "工具真实返回集合": [f'{r.title} {r.article_no}' for r in REGISTRY.refs],
        "数量": len(REGISTRY.refs),
    })

    # ③.b 候选边：登记一部图谱里没有、但本次真实命中的法（待人工确认）
    graph_law_titles = {n.name for n in KG.nodes.values() if n.type == "law"}
    candidates = []
    for r in REGISTRY.refs:
        if r.title not in graph_law_titles:
            eid = add_candidate_direct(case_name, r.title)
            if eid:
                candidates.append({"law": r.title, "edge_id": eid, "confirmed": False})
    step("STEP 3b · 图谱自进化：登记候选边（待人工确认）", {
        "候选边": candidates,
        "机制": "置信度 0.6 < 阈值 0.8 → 候选边；人工确认后才作为后续规划依据",
    } if candidates else {"候选边": [], "说明": "本次命中的法均已存在于图谱"})

    # ---------------- ④ 起草员：起草（故意埋雷） ----------------
    tmpl = await call(template_load, {"title": "合同审查意见书"})
    cite_lines = "\n".join(
        f"- {r.title} {r.article_no}；来源：{r.source}" for r in REGISTRY.refs)

    draft = f"""一、合同基本情况
本案为中外合资经营合同纠纷，标的为中日合资公司的出资义务履行与经营决策争议。
任务类型：{kind_meta(TASK_KIND)['label']}（{kind_meta(TASK_KIND)['goal']}）。
适用文书骨架：{'；'.join(tmpl['sections'])}。

二、逐条风险与修改建议
风险一（高）｜原文条款：出资时间表与违约责任约定不明。
  问题：未约定迟延出资的违约金计算方式与催告程序。
  修改建议：补充书面催告 + 逾期违约金条款，并明确继续履行请求权。
风险二（中）｜原文条款：董事会表决机制。
  问题：未约定僵局打破机制。
  修改建议：引入递进表决或第三方调解条款。

三、重点条款专项意见
1. 违约责任：依 {REGISTRY.refs[0].title if REGISTRY.refs else ''} 主张继续履行与损害赔偿；
2. 公司治理：依 {REGISTRY.refs[1].article_no if len(REGISTRY.refs) > 1 else ''} 的诚信履行原则，
   董事会决议应符合章程与股东会授权范围；
3. 争议解决：建议明确准据法与仲裁地。

**第9999条**（起草员另行援引，未在本次检索范围内 —— 这是一处故意埋入的库外条号）

四、引用自检
（待质检员填写）

五、总体结论与谈判策略
建议以继续履行为主、损害赔偿为辅的谈判基调，并在补充协议中固定出资时间表。
"""
    step("STEP 4 · 起草员起草（★ 故意埋入一处库外条号「第9999条」）", draft)

    # ---------------- ⑤ 质检员：两段自检 ----------------
    a1 = audit_citations(draft, REGISTRY.refs)
    first = {"stage": "第一段 · 引用是否真实存在",
             "ok": a1.ok, "total": a1.total, "covered": a1.covered,
             "unmatched": a1.unmatched, "report": audit_report(a1)}

    a2 = audit_applicability(REGISTRY.refs, sk["lawNames"])
    second = {"stage": "第二段 · 是否落在案由图谱请求权基础范围内",
              **a2.to_dict(), "report": applicability_report(a2)}

    step("STEP 5 · 质检员两段确定性自检（audit_citations + audit_applicability）", {
        "第一段_引用真实性": first,
        "第二段_法条适用性": second,
    })
    print(f"\n** 第一段结论：{a1.covered}/{a1.total} 命中，未覆盖 {a1.unmatched}")
    print(f"** 第二段结论：{applicability_report(a2).splitlines()[-1]}")

    # ---------------- ⑥ 交付 ----------------
    audit_block = (audit_report(a1) + "\n\n"
                   "**第二段：法条适用性校验（知识图谱）**\n" + applicability_report(a2))
    final = draft.replace("四、引用自检\n（待质检员填写）",
                          "四、引用自检\n" + audit_block)

    ok, note = export_docx(final, a1, a2, sk)
    step("STEP 6 · 交付", {
        "交付物": str(DELIVERABLE) if ok else "（未生成）",
        "导出说明": note,
        "第一段通过": a1.ok,
        "第二段超范围": a2.out_of_scope,
        "未覆盖项": a1.unmatched,
        "候选边（待人工确认）": [c["law"] for c in candidates],
    })

    log.update({
        "finished_at": time.strftime("%Y-%m-%dT%H:%M:%S"),
        "case": case_name,
        "skeleton": sk,
        "returned_refs": [r.to_dict() for r in REGISTRY.refs],
        "audit_stage1": first,
        "audit_stage2": second,
        "candidate_edges": candidates,
        "deliverable": str(DELIVERABLE) if ok else None,
    })
    RUNLOG.write_text(json.dumps(log, ensure_ascii=False, indent=2), encoding="utf-8")

    print(f"\n{'=' * 76}")
    print(f"取证日志：{RUNLOG}")
    if ok:
        print(f"交付物  ：{DELIVERABLE}")
    print("=" * 76)
    return 0


def export_docx(text: str, a1, a2, sk) -> tuple:
    """导出交付物 .docx。优先 python-docx；缺失时回退 Markdown。"""
    try:
        from docx import Document
    except Exception as exc:  # noqa: BLE001
        md = DELIVERABLE.with_suffix(".md")
        md.write_text(text, encoding="utf-8")
        return False, f"python-docx 不可用（{exc}），已改为导出 Markdown：{md}"

    doc = Document()
    doc.add_heading("合同审查意见书", level=0)
    doc.add_paragraph("法元 · Faron 文书智能体（小元） 出品")
    doc.add_paragraph(f"任务类型：{kind_meta(TASK_KIND)['label']}")
    doc.add_paragraph(f"命中案由：{sk['case']}")
    doc.add_paragraph(f"生成时间：{time.strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_paragraph("")

    for line in text.splitlines():
        s = line.rstrip()
        if not s:
            continue
        if len(s) > 1 and s[0] in "一二三四五" and s[1:2] == "、":
            doc.add_heading(s, level=1)
        elif s.startswith("- "):
            doc.add_paragraph(s[2:], style="List Bullet")
        elif s.startswith("**") and s.endswith("**"):
            doc.add_paragraph(s.strip("*"))
        else:
            doc.add_paragraph(s)

    # 第一段自检明细（Word 原生表格）
    doc.add_heading("附一：第一段引用自检明细（Word 原生表格）", level=1)
    t1 = doc.add_table(rows=1, cols=3)
    t1.style = "Light Grid Accent 1"
    for i, h in enumerate(["引用 token", "是否命中工具真实返回", "判定"]):
        t1.rows[0].cells[i].text = h
    for tok in (["—"] if a1.total == 0 else []):
        t1.add_row().cells[0].text = tok
    # 命中的与未命中的都以 token 形式记录
    for tok in _covered_tokens(text, a1):
        row = t1.add_row().cells
        row[0].text = tok
        row[1].text = "是"
        row[2].text = "有效引用"
    for tok in a1.unmatched:
        row = t1.add_row().cells
        row[0].text = tok
        row[1].text = "否"
        row[2].text = "未覆盖·不视为有效引用"

    # 第二段自检明细
    doc.add_heading("附二：第二段法条适用性校验", level=1)
    t2 = doc.add_table(rows=1, cols=3)
    t2.style = "Light Grid Accent 1"
    for i, h in enumerate(["图谱请求权基础", "本次引用法条", "判定"]):
        t2.rows[0].cells[i].text = h
    row = t2.add_row().cells
    row[0].text = "、".join(a2.law_names)
    row[1].text = "、".join(a2.in_scope + a2.out_of_scope)
    row[2].text = (f"范围内 {len(a2.in_scope)} 部"
                   + (f"；超范围 {len(a2.out_of_scope)} 部（需人工判断）"
                      if a2.out_of_scope else "；全部在范围内"))

    doc.save(str(DELIVERABLE))
    return True, f"已导出 .docx（{DELIVERABLE.stat().st_size} 字节）"


def _covered_tokens(text: str, a1) -> list:
    """命中的 token = 全部 token 减去未命中的。"""
    from faron_agentgroup import extract_citation_tokens
    return [t for t in extract_citation_tokens(text) if t not in a1.unmatched]


if __name__ == "__main__":
    sys.exit(asyncio.run(main()))
